"""
Per-market monthly signals from the analyst's v4 model, for the 2026-10-26 checks of his ideas inside the GRU
(read by src/data.py attach_checks through features.extra_flat_path / extra_flat_columns).

    python src/build_analyst_signals.py --workbook data/external/analyst_v4/20261021_Analyst_Maize_Forecast_v4_White.xlsx \
        --out data/external/20261026_Analyst_Signals_Monthly_v1.csv

Computed by HIS code (vendor/analyst_v4, via src/analyst_forecast.run_his_model), white maize series, origin month o:
  seasrel1/3/6   his "seasonal vs national": this market's usual seasonal move to o+h minus the national one (his
                 seasonal index, vintage = the year before o); origins from Jan 2016 (his first origin)
  wy_gap         white vs yellow maize in the same market, minus its mean over the previous 12 months
  sorg_gap       maize vs white sorghum in the same market, minus its previous-12-month mean
  wk_nat_slope   national: last weekly log price minus the first, inside month o (his weekly table: FEWS NET plus Kasuwa
  wk_nat_lw      national: last week minus the month's weekly mean                   where FEWS NET has no quote)
  wk_own_lw      this market's last week vs month mean, minus the national figure
  wk_own_l2      this market's last 2 weeks vs month mean, minus the national figure
  ip_anchor      his import-parity anchor: median of the last 24 months' gap to parity minus this month's gap
  fx_premium     street rate / official rate - 1 (his In_Macro)
Computed here, as he describes them in Feature_Tests (he tested and dropped each):
  duty_cut6      1 if import duty fell in the last 6 months (NADIH parity CSV duty column)
  duty_chg12     duty now minus duty 12 months earlier
  prem_alert     1 if fx_premium >= 0.40 (his alert threshold)
  price_vol3     s.d. of this market's last 3 monthly log price changes (our gap-filled panel)
  parity_vol3    s.d. of the last 3 monthly log changes of import parity
  ev_signed_3m   events from his In_Events tab that STARTED in months o-2..o, +1 if marked "up" for maize, -1 if "down",
                 national events for every market, regional ones for that region's markets. Start months only (an
                 event's end is hindsight). The list was compiled after the fact, so it is not a clean real-time feed.
Each signal has a mask column (1 = value present); missing values become 0 in the model with mask 0.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent))
from analyst_forecast import run_his_model  # noqa: E402

PARITY = "data/external/20261012_NADIH_ImportParity_Maize_Monthly_2006-01_2026-09_v1.csv"
PANEL = "data/panel_monthly_gapfilled_pit.parquet"


def his_signals(g: dict) -> pd.DataFrame:
    T, MK = g["MONTHS"], g["MK"]
    col = lambda tab, m: tab[("White", m)] if ("White", m) in tab.columns else pd.Series(np.nan, index=T)
    nat_slope, nat_lw, nat_l2 = g["NAT_SLOPE"], g["NAT_LW"], g["NAT_L2"]
    med_gap = g["MED24"] - g["GAP"]
    rows = []
    for m in MK:
        rows.append(pd.DataFrame({"market": m, "month": T, "wy_gap": col(g["SPRS"], m).values, "sorg_gap": col(g["RS_DEV"], m).values,
                                  "wk_nat_slope": nat_slope.reindex(T).values, "wk_nat_lw": nat_lw.reindex(T).values,
                                  "wk_own_lw": (col(g["LW"], m) - nat_lw).values, "wk_own_l2": (col(g["L2"], m) - nat_l2).values,
                                  "ip_anchor": col(med_gap, m).values, "fx_premium": g["PREM"].reindex(T).values}))
    d = pd.concat(rows, ignore_index=True)
    pn = g["pn"]
    w = pn[pn.type == "White"][["market", "origin", "f_seasrel1", "f_seasrel3", "f_seasrel6"]]
    w = w.rename(columns={"origin": "month", "f_seasrel1": "seasrel1", "f_seasrel3": "seasrel3", "f_seasrel6": "seasrel6"})
    return d.merge(w, on=["market", "month"], how="left")


def own_signals(d: pd.DataFrame, workbook: str) -> pd.DataFrame:
    q = pd.read_csv(PARITY)
    q["month"] = pd.to_datetime(q.month)
    duty = q.set_index("month").duty
    lp = np.log(q.set_index("month").parity_ngn_per_kg)
    t = pd.DataFrame({"duty_cut6": [float((duty.loc[:m].tail(7).diff() < 0).any()) for m in duty.index],
                      "duty_chg12": (duty - duty.shift(12)).values, "parity_vol3": lp.diff().rolling(3).std().values},
                     index=duty.index)
    d = d.merge(t, left_on="month", right_index=True, how="left")
    d["prem_alert"] = (d.fx_premium >= 0.40).astype(float).where(d.fx_premium.notna())
    pnl = pd.read_parquet(PANEL)
    pr = np.log(pnl.pivot(index="date", columns="market", values="price").sort_index())
    vol = pr.diff().rolling(3).std().stack().rename("price_vol3").reset_index().rename(columns={"date": "month"})
    d = d.merge(vol, on=["market", "month"], how="left")
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ev = [r for r in wb["In_Events"].iter_rows(min_row=5, values_only=True) if r[0] and str(r[0]).startswith("E")]
    zone = {r[0]: r[1] for r in wb["In_Zone_Bands"].iter_rows(min_row=5, max_row=40, min_col=4, max_col=5, values_only=True) if r[0]}
    region = {"North East": "NE", "North West": "NW", "North Central": "NC", "South West": "SW", "South South": "SS"}
    evs = pd.DataFrame([dict(start=pd.Timestamp(r[1] + "-01"), sign=1.0 if r[4] == "up" else (-1.0 if r[4] == "down" else 0.0),
                             region=r[5]) for r in ev])
    def score(m, mk):
        hit = evs[(evs.start <= m) & (evs.start > m - pd.DateOffset(months=3)) & evs.region.isin(["NAT", region.get(zone.get(mk), "")])]
        return float(hit.sign.sum())
    d["ev_signed_3m"] = [score(m, mk) for m, mk in zip(d.month, d.market)]
    return d


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    g = run_his_model(str(Path(a.workbook).resolve()))
    d = own_signals(his_signals(g), a.workbook)
    sig = ["seasrel1", "seasrel3", "seasrel6", "wy_gap", "sorg_gap", "wk_nat_slope", "wk_nat_lw", "wk_own_lw", "wk_own_l2",
           "ip_anchor", "fx_premium", "duty_cut6", "duty_chg12", "prem_alert", "price_vol3", "parity_vol3", "ev_signed_3m"]
    for s in sig:
        d[f"{s}_mask"] = d[s].notna().astype(float)
    d["origin_month"] = d.month.dt.strftime("%Y-%m")
    d[["market", "origin_month"] + sig + [f"{s}_mask" for s in sig]].to_csv(a.out, index=False)
    cov = {s: f"{int(d[s].notna().sum())} ({d.loc[d[s].notna(), 'origin_month'].min()}..{d.loc[d[s].notna(), 'origin_month'].max()})" for s in sig}
    print(f"{len(d)} market-months -> {a.out}")
    for k, v in cov.items():
        print(f"  {k:14s} {v}")


if __name__ == "__main__":
    main()
