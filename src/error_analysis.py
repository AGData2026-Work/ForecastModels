"""
Per-market error analysis workbook.

Reads every completed run under outputs/ plus the source panel, and writes one
xlsx with per-market MAE and MAPE, a recent predicted-versus-actual grid, an
aggregate summary on its own sheet, and a cleaning log.

    python src/error_analysis.py
    python src/error_analysis.py --since 2023-10-01
    python src/error_analysis.py --last-origins 8
    python src/error_analysis.py --arm conditional --out custom_name.xlsx

Recent window
-------------
Two definitions are produced because they answer different questions:

  Recent_By_Target   forecasts whose TARGET week falls on or after --since.
                     "How did we do on the most recent actual prices?"
                     Note this mixes forecast dates: a 26-week forecast landing
                     in September 2024 was issued in March 2024, while a 4-week
                     forecast landing the same week was issued in August.

  Recent_By_Origin   the last --last-origins forecast dates per market.
                     "How did we do on our most recent forecasts?"
                     Balanced across markets; each market contributes the same
                     number of origins, which matters because the origin grid is
                     staggered and each market's latest origin differs.

Proxy flags
-----------
Every row carries is_proxy_origin and is_proxy_target, joined from the panel's
is_proxy_price column. A large error against a proxy-filled actual is a data
question before it is a model question.
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HORIZONS = [4, 13, 26]
FONT = "Arial"


# ---------------------------------------------------------------- load
def load_runs(root: Path) -> tuple[pd.DataFrame, list[dict]]:
    frames, manifest = [], []
    for p in sorted(root.glob("*/*/predictions_paired.csv")):
        meta_path = p.parent / "run_metadata.json"
        meta = json.loads(meta_path.read_text()) if meta_path.exists() else {}
        df = pd.read_csv(p, parse_dates=["origin"])
        kind = meta.get("kind") or next(c for c in df.columns if c in ("RNN", "GRU"))
        build = meta.get("build", p.parent.parent.name)
        conv = meta.get("convention", "unknown")
        smoke = bool(meta.get("smoke", False))
        df = df.rename(columns={kind: "pred"})
        df["model"] = kind
        df["build"] = build
        df["convention"] = conv
        df["smoke"] = smoke
        df["run"] = f"{build}|{kind}|{conv}"
        frames.append(df)
        manifest.append(dict(run=df["run"].iloc[0], smoke=smoke, rows=len(df),
                             n_params=meta.get("n_params"),
                             seeds=str(meta.get("seeds")),
                             pairs_per_horizon=int(len(df) / df.h.nunique())))
    if not frames:
        raise SystemExit(f"no completed runs under {root}")
    return pd.concat(frames, ignore_index=True), manifest


def attach_panel(df: pd.DataFrame, panel_path: str) -> tuple[pd.DataFrame, dict]:
    """Add target week, proxy flags at origin and target, and market metadata."""
    d = pd.read_parquet(panel_path)
    df = df.copy()
    df["target"] = df.apply(
        lambda r: r["origin"] + pd.Timedelta(weeks=int(r["h"])), axis=1)

    px = d[["market", "date", "is_proxy_price"]]
    df = df.merge(px.rename(columns={"date": "origin",
                                     "is_proxy_price": "is_proxy_origin"}),
                  on=["market", "origin"], how="left")
    df = df.merge(px.rename(columns={"date": "target",
                                     "is_proxy_price": "is_proxy_target"}),
                  on=["market", "target"], how="left")
    meta = d.groupby("market").agg(state=("state", "first"),
                                   market_kind=("market_kind", "first"),
                                   upstream=("upstream_market", "first")).reset_index()
    df = df.merge(meta, on="market", how="left")

    log = {
        "rows": int(len(df)),
        "target_week_missing_from_panel": int(df["is_proxy_target"].isna().sum()),
        "origin_week_missing_from_panel": int(df["is_proxy_origin"].isna().sum()),
        "actuals_that_are_proxy_filled": int(df["is_proxy_target"].fillna(False).sum()),
        "origins_that_are_proxy_filled": int(df["is_proxy_origin"].fillna(False).sum()),
    }
    return df, log


# ---------------------------------------------------------------- metrics
def mae(a, p):
    return float(np.mean(np.abs(np.asarray(a, float) - np.asarray(p, float))))


def mape(a, p):
    a, p = np.asarray(a, float), np.asarray(p, float)
    ok = np.abs(a) > 1e-9
    return float(np.mean(np.abs((a[ok] - p[ok]) / a[ok])) * 100)


def bias(a, p):
    """Mean signed error. Positive means the model forecasts too high."""
    return float(np.mean(np.asarray(p, float) - np.asarray(a, float)))


def direction(a, p, p0):
    da = np.sign(np.asarray(a, float) - np.asarray(p0, float))
    dp = np.sign(np.asarray(p, float) - np.asarray(p0, float))
    ok = da != 0
    return float(np.mean(da[ok] == dp[ok]) * 100) if ok.any() else np.nan


def per_market(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (run, build, model, conv, mkt, h), g in df.groupby(
            ["run", "build", "model", "convention", "market", "h"]):
        base = dict(build=build, model=model, convention=conv, market=mkt,
                    state=g["state"].iloc[0], market_kind=g["market_kind"].iloc[0],
                    h_weeks=int(h), n=len(g),
                    mean_actual_NGN_per_kg=float(g["actual"].mean()))
        for label, col in (("challenger", "pred"), ("panel_fe", "panel_fe"),
                           ("naive", "naive")):
            base[f"{label}_MAE"] = mae(g["actual"], g[col])
            base[f"{label}_MAPE"] = mape(g["actual"], g[col])
        base["challenger_bias"] = bias(g["actual"], g["pred"])
        base["challenger_direction_pct"] = direction(g["actual"], g["pred"],
                                                     g["origin_price"])
        base["vs_panel_fe_MAE_pct"] = ((base["panel_fe_MAE"] - base["challenger_MAE"])
                                       / base["panel_fe_MAE"] * 100)
        base["vs_naive_MAE_pct"] = ((base["naive_MAE"] - base["challenger_MAE"])
                                    / base["naive_MAE"] * 100)
        base["beats_panel_fe"] = base["challenger_MAE"] < base["panel_fe_MAE"]
        base["proxy_actuals_in_cell"] = int(g["is_proxy_target"].fillna(False).sum())
        rows.append(base)
    out = pd.DataFrame(rows)
    return out.sort_values(["build", "convention", "model", "h_weeks",
                            "challenger_MAE"]).reset_index(drop=True)


def aggregate(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (run, build, model, conv, h), g in df.groupby(
            ["run", "build", "model", "convention", "h"]):
        r = dict(build=build, model=model, convention=conv, h_weeks=int(h), n=len(g),
                 challenger_MAE=mae(g["actual"], g["pred"]),
                 challenger_MAPE=mape(g["actual"], g["pred"]),
                 challenger_bias=bias(g["actual"], g["pred"]),
                 challenger_direction_pct=direction(g["actual"], g["pred"], g["origin_price"]),
                 panel_fe_MAE=mae(g["actual"], g["panel_fe"]),
                 panel_fe_MAPE=mape(g["actual"], g["panel_fe"]),
                 naive_MAE=mae(g["actual"], g["naive"]),
                 naive_MAPE=mape(g["actual"], g["naive"]),
                 markets_beating_panel_fe=None)
        pm = per_market(g)
        r["markets_beating_panel_fe"] = f"{int(pm.beats_panel_fe.sum())} of {len(pm)}"
        r["vs_panel_fe_MAE_pct"] = (r["panel_fe_MAE"] - r["challenger_MAE"]) / r["panel_fe_MAE"] * 100
        r["vs_naive_MAE_pct"] = (r["naive_MAE"] - r["challenger_MAE"]) / r["naive_MAE"] * 100
        rows.append(r)
    return pd.DataFrame(rows).sort_values(
        ["build", "convention", "model", "h_weeks"]).reset_index(drop=True)


def detail(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["error_NGN_per_kg"] = out["pred"] - out["actual"]
    out["abs_pct_error"] = (out["error_NGN_per_kg"] / out["actual"]).abs() * 100
    out["panel_fe_error"] = out["panel_fe"] - out["actual"]
    out["naive_error"] = out["naive"] - out["actual"]
    out["challenger_closer_than_panel_fe"] = (
        out["error_NGN_per_kg"].abs() < out["panel_fe_error"].abs())
    cols = ["build", "model", "convention", "market", "state", "h_weeks",
            "origin", "target", "origin_price", "actual", "pred",
            "error_NGN_per_kg", "abs_pct_error", "panel_fe", "panel_fe_error",
            "naive", "challenger_closer_than_panel_fe",
            "is_proxy_origin", "is_proxy_target", "pred_sd"]
    out = out.rename(columns={"h": "h_weeks"})
    have = [c for c in cols if c in out.columns]
    return out[have].sort_values(
        ["build", "convention", "model", "market", "h_weeks", "origin"]).reset_index(drop=True)


# ---------------------------------------------------------------- writing
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_sheet(wb: Workbook, name: str, df: pd.DataFrame, notes: list[str] | None = None,
                pct_cols: tuple[str, ...] = (), num_cols: tuple[str, ...] = ()) -> None:
    ws = wb.create_sheet(name[:31])
    r = 1
    if notes:
        for n in notes:
            ws.cell(row=r, column=1, value=n).font = Font(name=FONT, size=10, italic=True)
            r += 1
        r += 1
    header_row = r
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=j, value=c)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for i, (_, row) in enumerate(df.iterrows(), start=header_row + 1):
        for j, c in enumerate(df.columns, start=1):
            v = row[c]
            if isinstance(v, (pd.Timestamp, datetime)):
                v = v.date()
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = None if pd.isna(v) else float(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if c in pct_cols:
                cell.number_format = "0.00"
            elif c in num_cols:
                cell.number_format = "#,##0.00"
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for j, c in enumerate(df.columns, start=1):
        width = max(len(str(c)) + 2, 11)
        if c in ("market", "state", "build", "convention", "upstream", "market_kind"):
            width = max(width, 20)
        ws.column_dimensions[get_column_letter(j)].width = min(width, 30)
    ws.auto_filter.ref = (f"A{header_row}:"
                          f"{get_column_letter(len(df.columns))}{header_row + len(df)}")


def write_readme(wb: Workbook, args, manifest, clean_log, counts) -> None:
    ws = wb.create_sheet("README_and_Decisions", 0)
    lines = [
        ("Per-market error analysis: Nigerian maize price forecasting, Phase 2", True),
        ("", False),
        (f"Generated {datetime.now():%Y-%m-%d %H:%M}", False),
        ("All prices and errors in NGN per kilogram, exactly as reported in the source panel. "
         "No unit conversion applied.", False),
        ("", False),
        ("SHEETS", True),
        ("Summary_By_Horizon   aggregate metrics, one row per run and horizon. Kept on its own "
         "sheet, never beside the data.", False),
        ("Per_Market           MAE, MAPE, bias and directional accuracy per market, horizon and run.", False),
        ("Recent_By_Target     every forecast whose TARGET week falls on or after "
         f"{args.since}. Answers: how did we do on the most recent actual prices?", False),
        ("Recent_By_Origin     the last " + str(args.last_origins) + " forecast dates per market. "
         "Answers: how did we do on our most recent forecasts?", False),
        ("All_Forecasts        the full row-level detail, every run, every pair.", False),
        ("Run_Manifest         which runs were read, and whether any was a smoke run.", False),
        ("Cleaning_Log         joins, proxy counts and anything that needed checking.", False),
        ("", False),
        ("HOW TO READ THE TWO RECENT VIEWS", True),
        ("Recent_By_Target mixes forecast dates. A 26-week forecast landing in September 2024 was "
         "issued in March 2024; a 4-week forecast landing the same week was issued in August. So "
         "a 26-week row is not a later forecast, it is an older one reaching further.", False),
        ("Recent_By_Origin holds the number of origins per market constant. The origin grid is "
         "staggered, so each market's latest forecast date differs by up to three weeks; without "
         "this balancing, markets contribute unequally.", False),
        ("", False),
        ("COLUMN NOTES", True),
        ("challenger_bias      mean signed error. Positive means the model forecasts too high. "
         "A large bias with a modest MAE means a correctable level offset; a small bias with a "
         "large MAE means noise, which is not correctable by recalibration.", False),
        ("challenger_direction_pct   share of forecasts calling the direction of movement from the "
         "origin price correctly. 50 percent is a coin flip. Read alongside MAE: a model can lower "
         "MAE by predicting flat, which sinks this figure.", False),
        ("is_proxy_origin, is_proxy_target   TRUE where the panel's price for that week was itself "
         "proxy-filled rather than observed. A large error against a proxy actual is a data "
         "question before it is a model question.", False),
        ("vs_panel_fe_MAE_pct   positive means the challenger beat the incumbent.", False),
        ("pred_sd              spread of the prediction across random seeds. A market-level margin "
         "smaller than this is not a result.", False),
        ("", False),
        ("DECISIONS AND THINGS TO CROSSCHECK", True),
        ("1. Target weeks are computed as origin + h weeks. Confirmed against the source: the "
         "baseline grid uses exactly origin + 28, 91 and 182 days with no exceptions.", False),
        ("2. Actuals and origin prices come from the incumbent's own forecast file, not from the "
         "panel, so challenger and incumbent are scored against identical values.", False),
        ("3. Proxy flags are joined from the panel's is_proxy_price column on (market, week). "
         "A blank flag means that week is absent from the panel; the count is in Cleaning_Log.", False),
        ("4. MAPE divides by the actual price, so the low-price years of 2015 to 2019 carry far "
         "more weight than 2023 to 2024. MAE and MAPE can therefore rank the same markets "
         "differently. The change-control gate is set on MAE.", False),
        ("5. Per-market cell counts are small. At roughly 106 origins per market spread over three "
         "horizons, a single market and horizon holds about 106 comparisons; in the recent windows "
         "it holds far fewer. Treat per-market rankings as indicative, not decisive.", False),
        ("6. Any run flagged smoke in Run_Manifest covers a fraction of the grid and must not be "
         "compared against a full run.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, size=11, bold=bold)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 118
    r = len(lines) + 2
    ws.cell(row=r, column=1, value="ROW COUNTS").font = Font(name=FONT, size=11, bold=True)
    for k, v in counts.items():
        r += 1
        ws.cell(row=r, column=1, value=f"{k}: {v}").font = Font(name=FONT, size=11)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--outdir", default="outputs")
    ap.add_argument("--panel", default="data/panel_weekly.parquet")
    ap.add_argument("--since", default="2023-10-01",
                    help="Recent_By_Target includes targets on or after this date")
    ap.add_argument("--last-origins", type=int, default=8,
                    help="Recent_By_Origin keeps this many latest origins per market")
    ap.add_argument("--arm", default=None, choices=["unconditional", "conditional"],
                    help="restrict to one driver convention")
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    raw, manifest = load_runs(Path(a.outdir))
    if a.arm:
        raw = raw[raw["convention"] == a.arm]
        if not len(raw):
            raise SystemExit(f"no runs with convention={a.arm}")
    df, clean_log = attach_panel(raw, a.panel)

    since = pd.Timestamp(a.since)
    by_target = df[df["target"] >= since]
    if not len(by_target):
        latest = df["target"].max()
        fallback = latest - pd.Timedelta(weeks=52)
        print(f"WARNING --since {since.date()} selects no rows; the latest target in these "
              f"runs is {latest.date()}. Falling back to {fallback.date()} "
              f"(last 52 weeks of available targets).")
        since = fallback
        by_target = df[df["target"] >= since]

    keep = []
    for (run, mkt), g in df.groupby(["run", "market"]):
        latest = np.sort(g["origin"].unique())[-a.last_origins:]
        keep.append(g[g["origin"].isin(latest)])
    by_origin = pd.concat(keep, ignore_index=True) if keep else df.iloc[0:0]

    summary = aggregate(df)
    pm = per_market(df)
    pm_recent = per_market(by_target) if len(by_target) else pm.iloc[0:0]

    clean_log.update(
        recent_by_target_since=str(since.date()),
        recent_by_target_rows=int(len(by_target)),
        recent_by_target_window=(f"{by_target['target'].min().date()} to "
                                 f"{by_target['target'].max().date()}") if len(by_target) else "empty",
        recent_by_origin_last_n=a.last_origins,
        recent_by_origin_rows=int(len(by_origin)),
        latest_origin_per_market={k: str(pd.Timestamp(v).date()) for k, v in
                                  df.groupby("market")["origin"].max().items()},
        latest_target_in_data=str(df["target"].max().date()),
        smoke_runs=[m["run"] for m in manifest if m["smoke"]],
    )

    wb = Workbook()
    wb.remove(wb.active)

    write_sheet(wb, "Summary_By_Horizon", summary,
                notes=["Aggregate metrics per run and horizon. Errors in NGN/kg; MAPE in percent.",
                       "vs_panel_fe_MAE_pct positive means the challenger beat the incumbent."],
                num_cols=("challenger_MAE", "challenger_bias", "panel_fe_MAE", "naive_MAE"),
                pct_cols=("challenger_MAPE", "panel_fe_MAPE", "naive_MAPE",
                          "challenger_direction_pct", "vs_panel_fe_MAE_pct", "vs_naive_MAE_pct"))

    write_sheet(wb, "Per_Market", pm,
                notes=["Full evaluation period. Errors in NGN/kg; MAPE and percentages in percent.",
                       "Sorted by challenger MAE within each run and horizon: worst markets at the bottom."],
                num_cols=("mean_actual_NGN_per_kg", "challenger_MAE", "panel_fe_MAE",
                          "naive_MAE", "challenger_bias"),
                pct_cols=("challenger_MAPE", "panel_fe_MAPE", "naive_MAPE",
                          "challenger_direction_pct", "vs_panel_fe_MAE_pct", "vs_naive_MAE_pct"))

    if len(pm_recent):
        write_sheet(wb, "Per_Market_Recent", pm_recent,
                    notes=[f"Same layout, restricted to targets on or after {since.date()}.",
                           "Cell counts are small here. Check the n column before drawing conclusions."],
                    num_cols=("mean_actual_NGN_per_kg", "challenger_MAE", "panel_fe_MAE",
                              "naive_MAE", "challenger_bias"),
                    pct_cols=("challenger_MAPE", "panel_fe_MAPE", "naive_MAPE",
                              "challenger_direction_pct", "vs_panel_fe_MAE_pct", "vs_naive_MAE_pct"))

    write_sheet(wb, "Recent_By_Target", detail(by_target),
                notes=[f"Forecasts whose target week falls on or after {since.date()}.",
                       "Mixed forecast dates: a 26-week row was issued six months before its target.",
                       "challenger_closer_than_panel_fe compares absolute errors row by row."],
                num_cols=("origin_price", "actual", "pred", "error_NGN_per_kg",
                          "panel_fe", "panel_fe_error", "naive", "pred_sd"),
                pct_cols=("abs_pct_error",))

    write_sheet(wb, "Recent_By_Origin", detail(by_origin),
                notes=[f"The last {a.last_origins} forecast dates per market, balanced across markets.",
                       "Use this rather than Recent_By_Target when comparing markets to each other."],
                num_cols=("origin_price", "actual", "pred", "error_NGN_per_kg",
                          "panel_fe", "panel_fe_error", "naive", "pred_sd"),
                pct_cols=("abs_pct_error",))

    write_sheet(wb, "All_Forecasts", detail(df),
                notes=["Every scored forecast in every run. Row-level detail behind all other sheets."],
                num_cols=("origin_price", "actual", "pred", "error_NGN_per_kg",
                          "panel_fe", "panel_fe_error", "naive", "pred_sd"),
                pct_cols=("abs_pct_error",))

    write_sheet(wb, "Run_Manifest", pd.DataFrame(manifest),
                notes=["Runs read from disk. Any TRUE in the smoke column is a pipeline check, "
                       "not a result, and is not comparable to a full run."])

    cl = pd.DataFrame([{"item": k, "value": json.dumps(v) if isinstance(v, (dict, list)) else v}
                       for k, v in clean_log.items()])
    write_sheet(wb, "Cleaning_Log", cl,
                notes=["Every join, count and window definition behind this workbook."])

    write_readme(wb, a, manifest, clean_log, {
        "runs read": len(manifest),
        "total forecast rows": len(df),
        "per-market cells": len(pm),
        "Recent_By_Target rows": len(by_target),
        "Recent_By_Origin rows": len(by_origin),
    })

    stamp = datetime.now().strftime("%Y%m%d")
    out = a.out or f"outputs/{stamp}_Maize_Forecast_PerMarket_ErrorAnalysis_v1.xlsx"
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    print(f"runs read: {len(manifest)}")
    for m in manifest:
        flag = "  [SMOKE]" if m["smoke"] else ""
        print(f"  {m['run']}  {m['pairs_per_horizon']} pairs/horizon{flag}")
    print(f"\nlatest target in data: {df['target'].max().date()}")
    print(f"Recent_By_Target: {len(by_target)} rows, "
          f"{clean_log['recent_by_target_window']}")
    print(f"Recent_By_Origin: {len(by_origin)} rows, last {a.last_origins} origins per market")
    if clean_log["actuals_that_are_proxy_filled"]:
        print(f"NOTE {clean_log['actuals_that_are_proxy_filled']} scored actuals are "
              f"proxy-filled in the panel; flagged per row")
    if clean_log["smoke_runs"]:
        print(f"WARNING smoke runs present: {clean_log['smoke_runs']}")
    print(f"\nwrote {out}")


if __name__ == "__main__":
    main()
